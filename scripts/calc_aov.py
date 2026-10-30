# -*- coding: utf-8 -*-
"""客单价月度跟进：计算全部维度"""
import openpyxl, json
from datetime import datetime
from collections import defaultdict, OrderedDict

SRC = '/Users/SaciNa/Documents/臻选&芥子/臻选私域/三诺&华大成交用户清单.xlsx'
GROWTH = '/Users/SaciNa/WorkBuddy/2026-09-17-11-36-56/outputs/wbz_customer_growth_raw.json'
NAME_MAP = {"赵老师": "赵老师", "三诺健康顾问丨朱国银": "朱老师", "三诺健康管家 | 诺诺": "王老师",
            "三诺健康管家｜吴老师": "吴老师", "三诺健康管家 | 黎老师": "黎老师"}
SHORT = ["赵老师", "朱老师", "黎老师", "吴老师", "王老师"]
MONTHS = [5, 6, 7, 8, 9]

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb.active
# 按表头名定位列（源表曾插入「微信」列，硬编码 row[n] 会整体错位）
_hdr = [str(c).strip() if c is not None else '' for c in
        next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
def _col(name, fallback):
    return _hdr.index(name) if name in _hdr else fallback
IDX_NAME = _col('收件人', 4)
IDX_PHONE = _col('手机', 5)
rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[2] is None and row[3] is None:
        continue
    raw = row[1]
    if isinstance(raw, datetime):
        d = raw.date()
    else:
        s = str(raw).replace('下单时间：', '').strip()
        try:
            d = datetime.strptime(s, '%Y-%m-%d %H:%M:%S').date()
        except Exception:
            try:
                d = datetime.strptime(s, '%Y-%m-%d %H:%M').date()
            except Exception:
                d = None
    rows.append({'sales': str(row[0] or ''), 'date': d, 'product': str(row[2] or ''),
                 'amt': float(row[3] or 0),
                 'name': str(row[IDX_NAME] or ''), 'phone': str(row[IDX_PHONE] or '')})

first = {}
for r in sorted([x for x in rows if x['date']], key=lambda x: x['date']):
    if r['phone'] and r['phone'] not in first:
        first[r['phone']] = r['date'].month

# ── 月度客户基数（5月起累计净增）──
growth = json.load(open(GROWTH, encoding='utf-8'))
g = defaultdict(lambda: defaultdict(dict))
for r in growth:
    sh = NAME_MAP.get(r['staff_name'])
    if sh:
        g[r['_period']][sh] = {'add': r.get('increase_count', 0), 'lost': r.get('decrease_count', 0)}

def msub(m):
    return [r for r in rows if r['date'] and r['date'].month == m and r['date'].year == 2026]

cum = 0
monthly = []
for m in MONTHS:
    gg = g[f'2026-{m:02d}']
    add = sum(gg[s]['add'] for s in SHORT)
    lost = sum(gg[s]['lost'] for s in SHORT)
    cum += add - lost
    sub = msub(m)
    rev = sum(r['amt'] for r in sub)
    orders = len(sub)
    cust = len(set(r['phone'] for r in sub))
    grp = [r for r in sub if ('套组' in r['product'] or '+' in r['product'] or '合并' in r['product'])]
    grp_rev = sum(r['amt'] for r in grp)
    hi = [r for r in sub if r['amt'] >= 1000]
    monthly.append(dict(m=m, add=add, lost=lost, net=add - lost, cum=cum,
                        rev=rev, orders=orders, cust=cust,
                        aov_order=rev / orders if orders else 0,
                        aov_cust=rev / cust if cust else 0,
                        grp_n=len(grp), grp_rev=grp_rev, grp_aov=grp_rev / len(grp) if grp else 0,
                        single_n=orders - len(grp), single_rev=rev - grp_rev,
                        single_aov=(rev - grp_rev) / (orders - len(grp)) if orders - len(grp) else 0,
                        hi_n=len(hi), hi_rev=sum(r['amt'] for r in hi)))

# ── 顾问 × 月度 客单价矩阵 ──
matrix = OrderedDict()
for s in SHORT:
    matrix[s] = {}
    for m in MONTHS:
        sub = [r for r in msub(m) if r['sales'] == s]
        rev = sum(r['amt'] for r in sub)
        matrix[s][m] = dict(orders=len(sub), rev=rev,
                            aov=rev / len(sub) if sub else 0)
# 全组月度合计
matrix['合计'] = {m: dict(orders=next(x for x in monthly if x['m'] == m)['orders'],
                          rev=next(x for x in monthly if x['m'] == m)['rev'],
                          aov=next(x for x in monthly if x['m'] == m)['aov_order']) for m in MONTHS}

# ── 客单价分档 ──
BANDS = [(0, 300, '<300'), (300, 600, '300-600'), (600, 1000, '600-1000'),
         (1000, 2000, '1000-2000'), (2000, 999999, '>=2000')]
bands = {}
for m in MONTHS:
    sub = msub(m)
    bands[m] = []
    for lo, hi_v, lab in BANDS:
        gg = [r for r in sub if lo <= r['amt'] < hi_v]
        bands[m].append((lab, len(gg), sum(r['amt'] for r in gg)))

# ── 客单价区间中位/最高 ──
for x in monthly:
    sub = sorted(msub(x['m']), key=lambda r: r['amt'])
    amts = [r['amt'] for r in sub]
    x['max'] = amts[-1] if amts else 0
    x['median'] = amts[len(amts) // 2] if amts else 0

print(json.dumps({'monthly': monthly, 'bands': {str(k): v for k, v in bands.items()},
                  'matrix': {s: {str(m): v for m, v in d.items()} for s, d in matrix.items()}},
                 ensure_ascii=False, indent=1, default=str))
